"""Builds the investor data workbook (Lakeba-investor-data-*.xlsx).

  python3 build-investor-workbook.py out.xlsx blank
  python3 build-investor-workbook.py out.xlsx current export.json   # portal export (see docs)
  python3 build-investor-workbook.py out.xlsx rows rows.json        # any data: {"Companies": [[...], ...], ...}

In "rows" mode each sheet's rows are lists of values in column order (input
columns only); dates as "YYYY-MM-DD". Read back with scripts/import-investors.mjs.
"""
import json, sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.worksheet.formula import ArrayFormula

out, mode = sys.argv[1], sys.argv[2]  # mode: blank | current | rows
data = json.load(open(sys.argv[3])) if mode in ('current', 'rows') else None

F = 'Arial'
TITLE = Font(name=F, size=14, bold=True, color='1F2937')
DESC = Font(name=F, size=9, italic=True, color='6B7280')
HEAD = Font(name=F, size=10, bold=True, color='FFFFFF')
BODY = Font(name=F, size=10, color='0000FF')       # inputs: blue
CALC = Font(name=F, size=10, color='000000')       # formulas: black
EX = Font(name=F, size=10, italic=True, color='9CA3AF')
HEAD_IN = PatternFill('solid', fgColor='1E40AF')    # header over input column
HEAD_CALC = PatternFill('solid', fgColor='4B5563')  # header over check column
IN = PatternFill('solid', fgColor='FFF9DB')         # cells to fill in
GREY = PatternFill('solid', fgColor='F3F4F6')       # calculated
EXF = PatternFill('solid', fgColor='FAFAFA')
thin = Side(style='thin', color='E5E7EB')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
RED = PatternFill('solid', fgColor='FDE2E2')
REDF = Font(name=F, size=10, color='B91C1C', bold=True)

FIRST = 4      # first data row (row 2 headers, row 3 example)
ROWS = {'Companies': 60, 'Share classes': 200, 'Holders': 1000, 'People': 1000, 'Access': 1500, 'Holdings': 2500, 'Boards': 300, 'Prices': 500}
import os
if os.environ.get('SMALL'): ROWS = {k: min(v, 130) for k, v in ROWS.items()}
def last(sheet): return FIRST + ROWS[sheet] - 1
def rng(sheet, col): return f"'{sheet}'!${col}${FIRST}:${col}${last(sheet)}"

wb = Workbook()
start = wb.active
start.title = 'Start here'

LISTS = {
    'Relationship': ['Group', 'Controlled venture', 'Non-controlling investment'],
    'Portal access': ['Full portal', 'Portfolio entry only', 'Not in portal'],
    'Class type': ['Ordinary', 'Ordinary (non-voting)', 'Preference', 'Options', 'Convertible note', 'SAFE', 'Other'],
    'YesNo': ['Yes', 'No'],
    'Holder kind': ['Individual', 'Company', 'Trust', 'Super fund'],
    'Role': ['Investor', 'Board', 'Admin'],
    'Access level': ['Owner', 'Viewer'],
    'Position': ['Director', 'Chair', 'Company secretary', 'Observer', 'Company admin'],
    'Price source': ['Capital raise', 'Secondary trade', 'Directors’ valuation'],
    'Holding source': ['Share register', 'Registry CSV', 'Old portal (WordPress)', 'Manual'],
}

def sheet(name, title, desc, cols, example):
    """cols: list of (header, width, kind, note) — kind: in | calc; returns ws."""
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = TITLE
    ws.cell(row=1, column=3, value=desc).font = DESC
    for i, (h, w, kind, note, *_) in enumerate(cols, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = HEAD; c.fill = HEAD_IN if kind == 'in' else HEAD_CALC
        c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BOX
        if note: c.comment = Comment(note, 'Lakeba portal')
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[2].height = 32
    for i, v in enumerate(example, 1):
        c = ws.cell(row=3, column=i, value=v); c.font = EX; c.fill = EXF; c.border = BOX
    ws.cell(row=3, column=len(cols) + 1, value='← example row, not imported').font = DESC
    for r in range(FIRST, last(name) + 1):
        for i, (h, w, kind, *_) in enumerate(cols, 1):
            c = ws.cell(row=r, column=i)
            c.fill = IN if kind == 'in' else GREY
            c.font = BODY if kind == 'in' else CALC
            c.border = BOX
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    ws.auto_filter.ref = f"A2:{ws.cell(row=2, column=len(cols)).column_letter}{last(name)}"
    return ws

def listref(name):
    i = list(LISTS).index(name) + 1
    col = wb['Lists'].cell(row=1, column=i).column_letter
    return f"=Lists!${col}$2:${col}${len(LISTS[name]) + 1}"

def validate(ws, col, name, required=False):
    dv = DataValidation(type='list', formula1=listref(name), allow_blank=not required, showErrorMessage=True,
                        errorTitle='Choose from the list', error=f'Pick one of: {", ".join(LISTS[name])}')
    ws.add_data_validation(dv); dv.add(f"{col}{FIRST}:{col}{last(ws.title)}")

def fmt(ws, col, number_format):
    for r in range(3, last(ws.title) + 1): ws[f"{col}{r}"].number_format = number_format

def formula(ws, col, template, array=False):
    for r in range(FIRST, last(ws.title) + 1):
        f = template.format(r=r)
        ws[f"{col}{r}"] = ArrayFormula(f"{col}{r}", f) if array else f

def flag_not_ok(ws, col):
    ws.conditional_formatting.add(f"{col}{FIRST}:{col}{last(ws.title)}",
        FormulaRule(formula=[f'AND({col}{FIRST}<>"",{col}{FIRST}<>"OK")'], fill=RED, font=REDF))

# Lists sheet first (validations refer to it)
lists = wb.create_sheet('Lists')
for i, (k, vals) in enumerate(LISTS.items(), 1):
    lists.cell(row=1, column=i, value=k).font = Font(name=F, bold=True)
    for j, v in enumerate(vals, 2): lists.cell(row=j, column=i, value=v).font = Font(name=F)
    lists.column_dimensions[lists.cell(row=1, column=i).column_letter].width = 26

D = 'dd/mm/yyyy'
H = 'Holdings'

# ---------------------------------------------------------------- Companies
co = sheet('Companies', 'Companies', 'Lakeba Group and every company whose shareholders the portal could serve. One row per company.', [
    ('Company code *', 12, 'in', 'Short unique code, capitals, e.g. LKB, EZF. Used on every other sheet.'),
    ('Legal name *', 30, 'in', 'As registered, e.g. Ezifin Pty Ltd'),
    ('Display name', 18, 'in', 'Name shown in the portal, e.g. Ezifin'),
    ('ACN / ABN', 16, 'in', ''),
    ('Country', 12, 'in', ''),
    ('Relationship *', 22, 'in', 'Group, Controlled venture or Non-controlling investment'),
    ('Portal access *', 20, 'in', 'Full portal: its shareholders sign in. Portfolio entry only: shown to group shareholders as a Lakeba investment.'),
    ('Lakeba stake (stated)', 12, 'in', 'Lakeba’s ownership as reported, e.g. 65% (type the % sign).'),
    ('Stake as at', 12, 'in', 'Date of the stated stake'),
    ('Website', 24, 'in', ''),
    ('Portal admin email(s)', 28, 'in', 'Who manages this company in the portal (comma separated).'),
    ('Notes', 28, 'in', ''),
    ('Lakeba % per register', 14, 'calc', 'Calculated: latest Holdings of holder LKB-GROUP ÷ all latest holdings in this company’s classes that count toward %. Should match the stated stake.'),
    ('Check', 26, 'calc', ''),
], ['EZF', 'Ezifin Pty Ltd', 'Ezifin', '123 456 789', 'Australia', 'Controlled venture', 'Full portal', 0.65, date(2026, 6, 30), 'https://ezifin.com', 'cfo@ezifin.com', ''])
validate(co, 'F', 'Relationship', True); validate(co, 'G', 'Portal access', True)
fmt(co, 'H', '0.00%'); fmt(co, 'I', D); fmt(co, 'M', '0.00%')
formula(co, 'M', f'=IF(OR(A{{r}}="",F{{r}}="Group"),"",IFERROR(SUMIFS({rng(H,"C")},{rng(H,"H")},A{{r}},{rng(H,"A")},"LKB-GROUP",{rng(H,"J")},"Yes",{rng(H,"K")},"Yes")/SUMIFS({rng(H,"C")},{rng(H,"H")},A{{r}},{rng(H,"J")},"Yes",{rng(H,"K")},"Yes"),""))')
formula(co, 'N', f'=IF(A{{r}}="","",IF(COUNTIF($A${FIRST}:$A${last("Companies")},A{{r}})>1,"Duplicate code",IF(OR(B{{r}}="",F{{r}}="",G{{r}}=""),"Missing required field",IF(AND(F{{r}}="Controlled venture",H{{r}}<>"",M{{r}}<>""),IF(ABS(H{{r}}-M{{r}})>0.0005,"Stake differs from register","OK"),"OK"))))')
flag_not_ok(co, 'N')

# ---------------------------------------------------------------- Share classes
sc = sheet('Share classes', 'Share classes', 'Every class of shares (or notes/options) per company. Holdings refer to the class code.', [
    ('Company code *', 12, 'in', 'From the Companies sheet'),
    ('Class code *', 14, 'in', 'Unique across all companies: company code + class, e.g. EZF-ORD, LKB-F'),
    ('Class name *', 28, 'in', 'e.g. Ordinary shares'),
    ('Type *', 20, 'in', ''),
    ('Voting *', 9, 'in', 'Do holders of this class vote at meetings?'),
    ('Counts toward % *', 11, 'in', 'Include in ownership percentages? Usually Yes for shares, No for options, notes and SAFEs.'),
    ('Official on issue', 15, 'in', 'Total on issue per the share register. Percentages use this when present.'),
    ('Official as at', 12, 'in', 'Date of the official figure'),
    ('Currency', 9, 'in', 'AUD unless stated'),
    ('Notes', 26, 'in', ''),
    ('Total in this file', 15, 'calc', 'Calculated: sum of each holder’s latest holding in this class (Holdings sheet).'),
    ('Difference', 13, 'calc', 'Official on issue − total in this file. Not zero means the register in this file is incomplete or wrong.'),
    ('Check', 26, 'calc', ''),
], ['EZF', 'EZF-ORD', 'Ordinary shares', 'Ordinary', 'Yes', 'Yes', 10_000_000, date(2026, 6, 30), 'AUD', ''])
validate(sc, 'D', 'Class type', True); validate(sc, 'E', 'YesNo', True); validate(sc, 'F', 'YesNo', True)
fmt(sc, 'G', '#,##0'); fmt(sc, 'H', D); fmt(sc, 'K', '#,##0'); fmt(sc, 'L', '#,##0;[Red]-#,##0;"-"')
formula(sc, 'K', f'=IF(B{{r}}="","",SUMIFS({rng(H,"C")},{rng(H,"B")},B{{r}},{rng(H,"J")},"Yes"))')
formula(sc, 'L', '=IF(OR(B{r}="",G{r}=""),"",G{r}-K{r})')
formula(sc, 'M', f'=IF(B{{r}}="","",IF(COUNTIF($B${FIRST}:$B${last("Share classes")},B{{r}})>1,"Duplicate class code",IF(COUNTIF({rng("Companies","A")},A{{r}})=0,"Unknown company",IF(OR(C{{r}}="",D{{r}}="",E{{r}}="",F{{r}}=""),"Missing required field",IF(AND(G{{r}}<>"",H{{r}}=""),"Add the official as-at date",IF(AND(L{{r}}<>"",L{{r}}<>0),"Doesn’t add up to official","OK"))))))')
flag_not_ok(sc, 'M')

# ---------------------------------------------------------------- Holders
ho = sheet('Holders', 'Holders', 'The legal owners as they appear on the registers: people, companies, trusts, super funds. One row per holder, used across all companies.', [
    ('Holder ID *', 13, 'in', 'Your unique reference. Keep existing IDs (WP-0001…); new holders e.g. H-0001. LKB-GROUP is Lakeba Group Limited itself.'),
    ('Holder name (as on register) *', 40, 'in', 'Exactly as on the share register, e.g. Smith Super Pty Ltd ATF Smith Super Fund'),
    ('Kind *', 13, 'in', ''),
    ('ACN / ABN / trust details', 24, 'in', ''),
    ('Notes', 28, 'in', ''),
    ('Holdings rows', 10, 'calc', 'Calculated: rows on the Holdings sheet'),
    ('Logins with access', 10, 'calc', 'Calculated: people on the Access sheet who can see this holder'),
    ('Check', 26, 'calc', ''),
], ['H-0001', 'Jane Citizen', 'Individual', '', ''])
validate(ho, 'C', 'Holder kind', True)
formula(ho, 'F', f'=IF(A{{r}}="","",COUNTIF({rng(H,"A")},A{{r}}))')
formula(ho, 'G', f'=IF(A{{r}}="","",COUNTIF({rng("Access","B")},A{{r}}))')
formula(ho, 'H', f'=IF(A{{r}}="","",IF(COUNTIF($A${FIRST}:$A${last("Holders")},A{{r}})>1,"Duplicate holder ID",IF(OR(B{{r}}="",C{{r}}=""),"Missing required field",IF(F{{r}}=0,"No holdings",IF(AND(G{{r}}=0,A{{r}}<>"LKB-GROUP"),"No login (ok if not invited)","OK")))))')
ho.conditional_formatting.add(f"H{FIRST}:H{last('Holders')}", FormulaRule(formula=[f'H{FIRST}="No login (ok if not invited)"'], fill=PatternFill('solid', fgColor='FEF3C7'), font=Font(name=F, size=10, color='92400E')))
flag_not_ok(ho, 'H')

# ---------------------------------------------------------------- People
pe = sheet('People', 'People (portal logins)', 'Everyone who signs in. One row per email address; one person can see many holders across many companies.', [
    ('Email *', 32, 'in', 'Their sign-in email. One account per email.'),
    ('First name', 14, 'in', ''),
    ('Last name', 16, 'in', ''),
    ('Phone', 16, 'in', ''),
    ('Portal role *', 11, 'in', 'Investor for shareholders. Board for directors of Lakeba Group (venture directors go on the Boards sheet). Admin for Lakeba staff who manage the portal.'),
    ('Email notifications', 12, 'in', 'Yes: they get emails about new documents, meetings and news.'),
    ('Notes', 28, 'in', ''),
    ('Holders they see', 10, 'calc', 'Calculated: rows on the Access sheet'),
    ('Companies they see', 30, 'calc', 'Calculated from Access + Holdings (needs Excel 2019 or later).'),
    ('Check', 24, 'calc', ''),
], ['jane@example.com', 'Jane', 'Citizen', '+61 400 000 000', 'Investor', 'Yes', ''])
validate(pe, 'E', 'Role', True); validate(pe, 'F', 'YesNo')
formula(pe, 'H', f'=IF(A{{r}}="","",COUNTIF({rng("Access","A")},A{{r}}))')
formula(pe, 'I', f'=IF(A{{r}}="","",_xlfn.TEXTJOIN(", ",TRUE,IF(({rng("Companies","A")}<>"")*(COUNTIFS({rng("Access","A")},A{{r}},{rng("Access","F")},"*"&{rng("Companies","A")}&"*")>0),{rng("Companies","A")},"")))', array=True)
formula(pe, 'J', f'=IF(A{{r}}="","",IF(COUNTIF($A${FIRST}:$A${last("People")},A{{r}})>1,"Duplicate email",IF(ISERROR(FIND("@",A{{r}})),"Not an email address",IF(E{{r}}="","Missing role",IF(AND(H{{r}}=0,E{{r}}="Investor"),"Investor without holders","OK")))))')
flag_not_ok(pe, 'J')

# ---------------------------------------------------------------- Access
ac = sheet('Access', 'Access (who sees which holder)', 'Links a login to the holders it may see. A person who holds via a trust and personally has two rows.', [
    ('Email *', 32, 'in', 'From the People sheet'),
    ('Holder ID *', 13, 'in', 'From the Holders sheet'),
    ('Access *', 10, 'in', 'Owner can vote and ask questions; Viewer is read-only (e.g. an accountant).'),
    ('Notes', 26, 'in', ''),
    ('Holder name', 34, 'calc', 'Calculated'),
    ('Companies (via holdings)', 22, 'calc', 'Calculated: companies this holder holds shares in, so what this person will see.'),
    ('Check', 24, 'calc', ''),
], ['jane@example.com', 'H-0001', 'Owner', ''])
validate(ac, 'C', 'Access level', True)
formula(ac, 'E', f'=IF(B{{r}}="","",IFERROR(INDEX({rng("Holders","B")},MATCH(B{{r}},{rng("Holders","A")},0)),""))')
formula(ac, 'F', f'=IF(B{{r}}="","",_xlfn.TEXTJOIN(", ",TRUE,IF(({rng("Companies","A")}<>"")*(COUNTIFS({rng(H,"A")},B{{r}},{rng(H,"H")},{rng("Companies","A")},{rng(H,"J")},"Yes",{rng(H,"C")},">0")>0),{rng("Companies","A")},"")))', array=True)
formula(ac, 'G', f'=IF(A{{r}}="","",IF(COUNTIF({rng("People","A")},A{{r}})=0,"Email not on People",IF(COUNTIF({rng("Holders","A")},B{{r}})=0,"Unknown holder ID",IF(COUNTIFS($A${FIRST}:$A${last("Access")},A{{r}},$B${FIRST}:$B${last("Access")},B{{r}})>1,"Duplicate link",IF(C{{r}}="","Missing access","OK")))))')
flag_not_ok(ac, 'G')

# ---------------------------------------------------------------- Holdings
hd = sheet('Holdings', 'Holdings', 'Shares held, per holder, class and date. Add a new row when a holding changes; keep old rows as history. The latest date counts.', [
    ('Holder ID *', 13, 'in', 'From the Holders sheet'),
    ('Class code *', 13, 'in', 'From the Share classes sheet, e.g. EZF-ORD'),
    ('Shares *', 14, 'in', 'Whole number. Use 0 when a holder sold out (keeps the history).'),
    ('As at *', 12, 'in', 'Date of this balance'),
    ('Register ref', 14, 'in', 'Holder number on that company’s register, if any'),
    ('Source', 18, 'in', ''),
    ('Notes', 24, 'in', ''),
    ('Company', 10, 'calc', 'Calculated from the class'),
    ('Holder name', 32, 'calc', 'Calculated'),
    ('Latest?', 8, 'calc', 'Calculated: Yes when this is the most recent row for this holder and class'),
    ('Counts toward %', 9, 'calc', 'Calculated from the class'),
    ('Check', 24, 'calc', ''),
], ['H-0001', 'EZF-ORD', 25_000, date(2026, 6, 30), '', 'Share register', ''])
validate(hd, 'F', 'Holding source')
fmt(hd, 'C', '#,##0'); fmt(hd, 'D', D)
formula(hd, 'H', f'=IF(B{{r}}="","",IFERROR(INDEX({rng("Share classes","A")},MATCH(B{{r}},{rng("Share classes","B")},0)),""))')
formula(hd, 'I', f'=IF(A{{r}}="","",IFERROR(INDEX({rng("Holders","B")},MATCH(A{{r}},{rng("Holders","A")},0)),""))')
formula(hd, 'J', f'=IF(OR(A{{r}}="",B{{r}}="",D{{r}}=""),"",IF(D{{r}}=_xlfn.MAXIFS($D${FIRST}:$D${last(H)},$A${FIRST}:$A${last(H)},A{{r}},$B${FIRST}:$B${last(H)},B{{r}}),"Yes","No"))')
formula(hd, 'K', f'=IF(B{{r}}="","",IFERROR(INDEX({rng("Share classes","F")},MATCH(B{{r}},{rng("Share classes","B")},0)),""))')
formula(hd, 'L', f'=IF(A{{r}}="","",IF(I{{r}}="","Unknown holder ID",IF(H{{r}}="","Unknown class code",IF(OR(C{{r}}="",D{{r}}=""),"Missing shares or date",IF(OR(C{{r}}<0,C{{r}}<>INT(C{{r}})),"Shares must be a whole number",IF(COUNTIFS($A${FIRST}:$A${last(H)},A{{r}},$B${FIRST}:$B${last(H)},B{{r}},$D${FIRST}:$D${last(H)},D{{r}})>1,"Duplicate row (same date)","OK"))))))')
flag_not_ok(hd, 'L')

# ---------------------------------------------------------------- Boards
bo = sheet('Boards', 'Boards and company admins', 'Directors and officers per company (including Lakeba Group). Gives board access to that company only.', [
    ('Company code *', 12, 'in', ''),
    ('Email *', 32, 'in', 'Must also be on the People sheet'),
    ('Position *', 18, 'in', 'Company admin: manages this company in the portal (not other companies).'),
    ('From', 12, 'in', ''),
    ('To', 12, 'in', 'Leave empty while current'),
    ('Notes', 26, 'in', ''),
    ('Check', 24, 'calc', ''),
], ['EZF', 'jane@example.com', 'Director', date(2024, 1, 1), None, ''])
validate(bo, 'C', 'Position', True); fmt(bo, 'D', D); fmt(bo, 'E', D)
formula(bo, 'G', f'=IF(A{{r}}="","",IF(COUNTIF({rng("Companies","A")},A{{r}})=0,"Unknown company",IF(COUNTIF({rng("People","A")},B{{r}})=0,"Email not on People",IF(C{{r}}="","Missing position","OK"))))')
flag_not_ok(bo, 'G')

# ---------------------------------------------------------------- Prices
pr = sheet('Prices', 'Reference prices', 'Indicative price per share, used for indicative values. One row per class and date.', [
    ('Class code *', 13, 'in', ''),
    ('Price per share *', 14, 'in', ''),
    ('Currency', 9, 'in', 'AUD unless stated'),
    ('Priced on *', 12, 'in', ''),
    ('Source *', 20, 'in', ''),
    ('Note', 34, 'in', 'e.g. Series A at $1.20'),
    ('Check', 24, 'calc', ''),
], ['EZF-ORD', 1.2, 'AUD', date(2026, 3, 31), 'Capital raise', 'Seed round'])
validate(pr, 'E', 'Price source', True); fmt(pr, 'B', '$#,##0.0000'); fmt(pr, 'D', D)
formula(pr, 'G', f'=IF(A{{r}}="","",IF(COUNTIF({rng("Share classes","B")},A{{r}})=0,"Unknown class code",IF(OR(B{{r}}="",D{{r}}="",E{{r}}=""),"Missing required field",IF(COUNTIFS($A${FIRST}:$A${last("Prices")},A{{r}},$D${FIRST}:$D${last("Prices")},D{{r}},$E${FIRST}:$E${last("Prices")},E{{r}})>1,"Duplicate price","OK"))))')
flag_not_ok(pr, 'G')

# ---------------------------------------------------------------- Checks
ch = wb.create_sheet('Checks', 1)
ch['A1'] = 'Checks'; ch['A1'].font = TITLE
ch['A2'] = 'Everything should read 0 before the file is imported. Filter a sheet’s Check column to find the rows.'; ch['A2'].font = DESC
ch.column_dimensions['A'].width = 26; ch.column_dimensions['B'].width = 14; ch.column_dimensions['C'].width = 14; ch.column_dimensions['D'].width = 60
for i, h in enumerate(['Sheet', 'Rows filled', 'Problems', 'What it checks'], 1):
    c = ch.cell(row=4, column=i, value=h); c.font = HEAD; c.fill = HEAD_CALC; c.border = BOX
checks = [
    ('Companies', 'A', 'N', 'Unique code, required fields, stated Lakeba stake matches the register'),
    ('Share classes', 'B', 'M', 'Unique code, known company, holdings add up to the official number on issue'),
    ('Holders', 'A', 'H', 'Unique ID, required fields, has holdings (“No login” is a warning, counted here too)'),
    ('People', 'A', 'J', 'Unique email, valid address, investors linked to at least one holder'),
    ('Access', 'A', 'G', 'Email and holder exist, no duplicate links'),
    ('Holdings', 'A', 'L', 'Holder and class exist, whole shares, no duplicate rows'),
    ('Boards', 'A', 'G', 'Company and email exist'),
    ('Prices', 'A', 'G', 'Class exists, no duplicates'),
]
for i, (s, key, chk, what) in enumerate(checks, 5):
    ch.cell(row=i, column=1, value=s).font = Font(name=F, size=10, bold=True)
    ch.cell(row=i, column=2, value=f'=SUMPRODUCT(({rng(s, key)}<>"")*1)').font = CALC
    c = ch.cell(row=i, column=3, value=f'=SUMPRODUCT(({rng(s, chk)}<>"")*({rng(s, chk)}<>"OK"))'); c.font = CALC
    ch.cell(row=i, column=4, value=what).font = Font(name=F, size=10)
    for col in range(1, 5): ch.cell(row=i, column=col).border = BOX
ch.conditional_formatting.add('C5:C12', FormulaRule(formula=['C5>0'], fill=RED, font=REDF))
r = 14
ch.cell(row=r, column=1, value='Ownership by company').font = Font(name=F, size=12, bold=True)
for i, h in enumerate(['Company', 'Holders (latest > 0)', 'Shares (count toward %)', 'Lakeba % per register'], 1):
    c = ch.cell(row=r + 1, column=i, value=h); c.font = HEAD; c.fill = HEAD_CALC; c.border = BOX
for k in range(12):
    rr = r + 2 + k
    ch.cell(row=rr, column=1, value=f"=IF('Companies'!A{FIRST + k}=\"\",\"\",'Companies'!A{FIRST + k})")
    ch.cell(row=rr, column=2, value=f'=IF(A{rr}="","",SUMPRODUCT(({rng(H,"H")}=A{rr})*({rng(H,"J")}="Yes")*({rng(H,"C")}>0)))')
    ch.cell(row=rr, column=3, value=f'=IF(A{rr}="","",SUMIFS({rng(H,"C")},{rng(H,"H")},A{rr},{rng(H,"J")},"Yes",{rng(H,"K")},"Yes"))').number_format = '#,##0'
    ch.cell(row=rr, column=4, value=f"=IF(A{rr}=\"\",\"\",'Companies'!M{FIRST + k})").number_format = '0.00%'
    for col in range(1, 5): ch.cell(row=rr, column=col).font = CALC; ch.cell(row=rr, column=col).border = BOX

# ---------------------------------------------------------------- Start here
start['A1'] = 'Lakeba Investor Portal: investor data'; start['A1'].font = Font(name=F, size=16, bold=True)
lines = [
    ('', ''),
    ('What this file is', 'One place to record every company, share class, holder, login and holding the portal should know about: Lakeba Group and its ventures. The portal is built and checked from it.'),
    ('Version', {'current': 'Pre-filled with the data in the portal today (Lakeba Group only). Check each row; add the ventures.', 'rows': 'Filled from a data file.', 'blank': 'Blank template.'}[mode]),
    ('', ''),
    ('How to fill it in', ''),
    ('1. Companies', 'Lakeba Group and each venture, with a short code (LKB, EZF…).'),
    ('2. Share classes', 'Each class per company. Class codes start with the company code (EZF-ORD). Add the official number on issue from the register.'),
    ('3. Holders', 'Every legal owner on any register, once, even if they hold in several companies. LKB-GROUP is Lakeba Group Limited (it holds the ventures).'),
    ('4. Holdings', 'One row per holder, class and date. To record a change, add a row with the new date; don’t overwrite.'),
    ('5. People', 'Everyone who will sign in, one row per email.'),
    ('6. Access', 'Which holders each person sees. What they see follows from the holdings: an Ezifin-only shareholder sees only Ezifin.'),
    ('7. Boards', 'Directors and company admins per company.'),
    ('8. Prices', 'Reference prices per class (optional).'),
    ('9. Checks', 'Every “Problems” figure should be 0. Red cells on each sheet show which rows.'),
    ('', ''),
    ('Legend', ''),
    ('Yellow cells, blue text', 'Fill these in. Columns marked * are required.'),
    ('Grey cells, black text', 'Calculated. Don’t type here.'),
    ('Light grey italic row 3', 'An example of the format. It is not imported.'),
    ('Red cell', 'A check found a problem in that row.'),
    ('Header notes', 'Hover over a column header for an explanation.'),
    ('', ''),
    ('Formats', 'Dates as dd/mm/yyyy. Shares as whole numbers. Percentages with a % sign (65%). Choose list values from the drop-downs.'),
    ('Confidential', 'This file contains personal information about investors. Keep it off email and shared drives that aren’t access-controlled.' if mode != 'blank' else 'Once filled in, this file contains personal information about investors. Keep it access-controlled.'),
]
for i, (a, b) in enumerate(lines, 2):
    start.cell(row=i, column=1, value=a).font = Font(name=F, size=10, bold=bool(b) or a in ('How to fill it in', 'Legend'))
    c = start.cell(row=i, column=2, value=b); c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
start.column_dimensions['A'].width = 26; start.column_dimensions['B'].width = 110
for a, fill, font in [(18, IN, BODY), (19, GREY, CALC), (20, EXF, EX), (21, RED, REDF)]:
    start.cell(row=a, column=1).fill = fill
lists.sheet_state = 'hidden'

# ---------------------------------------------------------------- Data
companies = [
    ('LKB', 'Lakeba Group Limited', 'Lakeba Group', 'Group', 'Full portal'),
    ('EZF', '', 'Ezifin', 'Controlled venture', 'Full portal'),
    ('DOX', '', 'DoxAI', 'Controlled venture', 'Full portal'),
    ('BRK', '', 'Bricklet', 'Controlled venture', 'Full portal'),
    ('SHF', '', 'Shelfie', 'Controlled venture', 'Full portal'),
    ('WHK', '', 'Whitehawk', 'Non-controlling investment', 'Portfolio entry only'),
    ('MIST', '', 'MIST', 'Non-controlling investment', 'Portfolio entry only'),
    ('AST', '', 'Assetora', 'Non-controlling investment', 'Portfolio entry only'),
    ('APR', '', 'Appreci', 'Non-controlling investment', 'Portfolio entry only'),
    ('ZUC', '', 'Zucoins', 'Non-controlling investment', 'Portfolio entry only'),
    ('EZP', '', 'EziProp', 'Non-controlling investment', 'Portfolio entry only'),
]
for i, (code, legal, disp, rel, acc) in enumerate(companies, FIRST):
    for col, v in zip('ABCFG', (code, legal or None, disp, rel, acc)): co[f'{col}{i}'] = v
    co[f'E{i}'] = 'Australia' if code == 'LKB' else None

classes = [('LKB', 'LKB-ORD', 'Ordinary shares', 'Ordinary', 'Yes', 'Yes'), ('LKB', 'LKB-F', 'F class shares (non-voting)', 'Ordinary (non-voting)', 'No', 'Yes')]
classes += [(c, f'{c}-ORD', 'Ordinary shares', 'Ordinary', 'Yes', 'Yes') for c in ('EZF', 'DOX', 'BRK', 'SHF')]
for i, row in enumerate(classes, FIRST):
    for col, v in zip('ABCDEF', row): sc[f'{col}{i}'] = v
    sc[f'I{i}'] = 'AUD'

ho[f'A{FIRST}'], ho[f'B{FIRST}'], ho[f'C{FIRST}'], ho[f'E{FIRST}'] = 'LKB-GROUP', 'Lakeba Group Limited', 'Company', 'Holds Lakeba’s shares in each venture'

if data and mode == 'current':
    kinds = {'individual': 'Individual', 'company': 'Company', 'trust': 'Trust', 'super_fund': 'Super fund'}
    ref = {h['id']: h['registry_ref'] for h in data['holders']}
    for i, h in enumerate(data['holders'], FIRST + 1):
        ho[f'A{i}'], ho[f'B{i}'], ho[f'C{i}'] = h['registry_ref'], h['name'], kinds[h['kind']]
    people = data['people']
    emails = {p['email'] for p in people}
    for i, p in enumerate(people, FIRST):
        pe[f'A{i}'], pe[f'B{i}'], pe[f'C{i}'], pe[f'D{i}'] = p['email'], p['first_name'] or None, p['last_name'] or None, p['phone']
        pe[f'E{i}'] = p['role'].capitalize(); pe[f'F{i}'] = 'Yes' if p['notify_email'] else 'No'
    acc = [a for a in data['access'] if a['email'] in emails]
    for i, a in enumerate(sorted(acc, key=lambda a: (a['email'], ref[a['holder_id']])), FIRST):
        ac[f'A{i}'], ac[f'B{i}'], ac[f'C{i}'] = a['email'], ref[a['holder_id']], a['access'].capitalize()
    for i, x in enumerate(data['holdings'], FIRST):
        hd[f'A{i}'], hd[f'B{i}'], hd[f'C{i}'] = ref[x['holder_id']], f"LKB-{x['code']}", x['shares']
        hd[f'D{i}'] = date.fromisoformat(x['as_at']); hd[f'F{i}'] = 'Old portal (WordPress)'
    src = {'capital_raise': 'Capital raise', 'secondary_trade': 'Secondary trade', 'valuation': 'Directors’ valuation'}
    for i, p in enumerate(data['prices'], FIRST):
        pr[f'A{i}'], pr[f'B{i}'], pr[f'C{i}'] = f"LKB-{p['code']}", float(p['price']), p['currency']
        pr[f'D{i}'], pr[f'E{i}'], pr[f'F{i}'] = date.fromisoformat(p['priced_on']), src[p['source']], p['note']
    for c in data['classes']:
        if c['official_on_issue']:
            r = FIRST + [x[1] for x in classes].index(f"LKB-{c['code']}")
            sc[f'G{r}'], sc[f'H{r}'] = c['official_on_issue'], date.fromisoformat(c['official_as_at'])

if mode == 'rows':
    import re
    def val(v):
        return date.fromisoformat(v) if isinstance(v, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', v) else v
    for name in ['Companies', 'Share classes', 'Holders', 'People', 'Access', 'Holdings', 'Boards', 'Prices']:
        ws = wb[name]
        for r in range(FIRST, last(name) + 1):           # replace the defaults
            for c in range(1, ws.max_column + 1):
                if ws.cell(row=2, column=c).fill.fgColor.rgb.endswith('1E40AF'): ws.cell(row=r, column=c).value = None
        for i, row in enumerate(data.get(name, []), FIRST):
            for j, v in enumerate(row, 1): ws.cell(row=i, column=j, value=val(v))

wb.move_sheet('Checks', offset=-(wb.index(wb['Checks']) - 1))
wb.calculation.fullCalcOnLoad = True
wb.save(out)
print('saved', out)
